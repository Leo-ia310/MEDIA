from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = r"C:\Users\Maykel\Documents\MEDIA\Informe_Feria_STEM_2026_MEDIA.docx"
LOGO = r"C:\Users\Maykel\Documents\MEDIA\docx_work\template_media\image1.png"

participants = [
    "Maikel Leandro Martinez Flores",
    "Michael Yurem Rugama Montenegro",
    "Osmar Alejandro Castro Mercado",
    "Lleyton Antonio Mendez Montiel",
    "Helmut",
]


def format_run(run, bold=False, italic=False, size=12, font="Arial", color=(0, 0, 0)):
    run.bold = bold
    run.italic = italic
    run.font.name = font
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(*color)
    run._element.rPr.rFonts.set(qn("w:ascii"), font)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), font)


def paragraph(doc, text="", bold=False, italic=False, align=None, size=12, after=6, before=0, first_line=False):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.space_before = Pt(before)
    if first_line:
        p.paragraph_format.first_line_indent = Inches(0.35)
    run = p.add_run(text)
    format_run(run, bold=bold, italic=italic, size=size)
    return p


def heading(doc, number, text):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(f"{number}.\t{text}")
    format_run(r, bold=True)
    return p


def bullet(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run("- " + text)
    format_run(r)
    return p


def set_cell_border(cell, color="D9D9D9", size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = f"w:{edge}"
        element = tc_borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            tc_borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)
        element.set(qn("w:space"), "0")


def set_cell_margins(cell, top=120, start=120, bottom=120, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def style_table(table, header_fill="1F4E79"):
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)
    for row_i, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_border(cell)
            set_cell_margins(cell)
            for p in cell.paragraphs:
                p.paragraph_format.line_spacing = 1.15
                p.paragraph_format.space_after = Pt(0)
                for run in p.runs:
                    format_run(run, size=10)
            if row_i == 0:
                shade_cell(cell, header_fill)
                for p in cell.paragraphs:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for run in p.runs:
                        run.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)
            elif row_i % 2 == 0:
                shade_cell(cell, "F3F6FA")


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    for i, title in enumerate(headers):
        table.rows[0].cells[i].text = title
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = value
    style_table(table)
    return table


def setup_document():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(12)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    return doc


def build():
    doc = setup_document()

    # Portada.
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(6)
    p.add_run().add_picture(LOGO, width=Inches(2.2))
    paragraph(doc, "TU UNIVERSIDAD POR EXCELENCIA", bold=True, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, size=12, after=8)

    box = doc.add_table(rows=1, cols=1)
    cell = box.cell(0, 0)
    set_cell_border(cell, color="000000", size="18")
    set_cell_margins(cell, top=120, bottom=120)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = cell.paragraphs[0].add_run("FACULTAD DE INGENIERÍAS Y ARQUITECTURA")
    format_run(r, bold=True, size=14)
    p2 = cell.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p2.add_run("Feria STEM 2026 - 4ta Edición")
    format_run(r, bold=True, size=13)

    paragraph(doc, "", after=6)
    paragraph(doc, "Nombre del proyecto:", bold=True, after=2)
    paragraph(doc, "MEDIA: tutor médico educativo con inteligencia artificial y biblioteca científica", after=8)
    paragraph(doc, "Participantes:", bold=True, after=2)
    for name in participants:
        paragraph(doc, f"{name} - Integrante del equipo desarrollador", after=0, size=10.5)
    paragraph(doc, "", after=6)
    paragraph(doc, "Tutor:", bold=True, after=2)
    paragraph(doc, "Pendiente de asignación institucional", after=0)
    paragraph(doc, "Fecha de presentación: Octubre, 2026", align=WD_ALIGN_PARAGRAPH.RIGHT, after=0, before=4)

    doc.add_page_break()
    paragraph(doc, "Índice", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=12)
    for item in [
        "I. Introducción",
        "II. Planteamiento de problema",
        "III. Justificación",
        "IV. Objetivos",
        "V. Desarrollo técnico y metodológico",
        "VI. Impacto del proyecto",
        "VII. Conclusiones",
        "VIII. Recomendaciones",
        "IX. Referencias bibliográficas",
        "X. Anexos",
    ]:
        paragraph(doc, item, after=2)

    doc.add_page_break()

    heading(doc, "I", "Introducción")
    paragraph(doc, "MEDIA es una plataforma web de apoyo educativo para estudiantes de medicina y ciencias de la salud. El proyecto integra una interfaz de chat con inteligencia artificial, autenticación de usuarios, historial de conversaciones, biblioteca de documentos médicos, generación de recursos de estudio y mecanismos de verificación para reducir respuestas sin respaldo documental.", first_line=True)
    paragraph(doc, "La importancia del proyecto se encuentra en la necesidad de contar con herramientas digitales que ayuden al estudiante a estudiar de forma ordenada, segura y basada en fuentes. En el área médica, una explicación incorrecta puede causar confusión académica o malas interpretaciones; por eso MEDIA no se plantea como sustituto de un profesional de salud, sino como un tutor educativo que organiza información y orienta el aprendizaje.", first_line=True)
    paragraph(doc, "La tecnología utilizada combina un frontend web desarrollado con HTML, CSS y JavaScript, un backend en FastAPI, almacenamiento y autenticación con Supabase, modelos de lenguaje mediante Groq, respaldo de texto con Cloudflare Workers AI, generación de imágenes educativas con Gemini y una arquitectura RAG preparada para consultar documentos médicos locales y cargados en Supabase.", first_line=True)

    heading(doc, "II", "Planteamiento de problema")
    paragraph(doc, "Los estudiantes de medicina suelen enfrentarse a grandes cantidades de información dispersa en libros, guías, artículos, apuntes y recursos digitales. Cuando se usan asistentes de inteligencia artificial sin control, existe el riesgo de recibir respuestas que parecen correctas pero que no citan fuentes verificables, mezclan datos o presentan afirmaciones médicas sin suficiente evidencia.", first_line=True)
    paragraph(doc, "El problema que busca atender MEDIA es la falta de una herramienta educativa centralizada que combine conversación guiada, memoria de estudio, biblioteca médica, recuperación de referencias y generación de materiales didácticos dentro de una misma experiencia. La necesidad no es reemplazar el juicio médico, sino mejorar el proceso de aprendizaje mediante explicaciones claras, trazabilidad de fuentes y recursos de práctica.", first_line=True)

    heading(doc, "III", "Justificación")
    paragraph(doc, "El equipo decidió realizar MEDIA porque la medicina requiere estudio constante, comprensión profunda y consulta responsable de fuentes. Una plataforma de tutoría con IA puede ayudar a repasar conceptos, generar preguntas de práctica, resumir temas, preparar presentaciones y orientar la lectura de documentos, siempre que incorpore límites de seguridad y controles contra la invención de citas.", first_line=True)
    paragraph(doc, "El proyecto sirve como apoyo académico para estudiantes de medicina, enfermería y ciencias de la salud. Beneficia directamente a quienes necesitan organizar su aprendizaje, revisar temas complejos, practicar con quizzes y acceder a referencias médicas desde una biblioteca digital. También beneficia a docentes o tutores, porque muestra una propuesta tecnológica adaptable al estudio supervisado.", first_line=True)
    paragraph(doc, "MEDIA tiene valor STEM porque une ingeniería de software, inteligencia artificial, bases de datos, diseño de interfaz, procesamiento de documentos y criterios de seguridad educativa. Además, demuestra cómo una solución digital puede responder a una necesidad real del entorno académico sin presentar diagnósticos definitivos ni fomentar automedicación.", first_line=True)

    heading(doc, "IV", "Objetivos")
    paragraph(doc, "Objetivo general:", bold=True, after=2)
    paragraph(doc, "Desarrollar una plataforma web educativa llamada MEDIA que utilice inteligencia artificial y una biblioteca de referencias médicas para apoyar el estudio de estudiantes de medicina y ciencias de la salud de forma clara, organizada y segura.", first_line=True)
    paragraph(doc, "Objetivos específicos:", bold=True, after=2)
    bullet(doc, "Diseñar una interfaz web tipo chat que permita iniciar conversaciones, guardar historial, buscar chats y acceder a herramientas educativas.")
    bullet(doc, "Implementar un backend con FastAPI capaz de gestionar autenticación, conversaciones, perfil de aprendizaje, biblioteca, retroalimentación y generación de recursos.")
    bullet(doc, "Integrar proveedores de inteligencia artificial para texto, razonamiento, visión e imágenes educativas, usando Groq como proveedor principal, Cloudflare como respaldo y Gemini para imágenes.")
    bullet(doc, "Preparar una arquitectura RAG con documentos médicos, fragmentos verificables y controles anti-alucinación para diferenciar respuestas verificadas, conocimiento no verificado e insuficiencia de evidencia.")

    heading(doc, "V", "Desarrollo técnico y metodológico")
    paragraph(doc, "El desarrollo de MEDIA se organizó como una aplicación web con separación entre frontend, backend, base de datos y servicios de inteligencia artificial. El frontend se construyó con archivos estáticos que contienen la interfaz, estilos, navegación lateral, chat, biblioteca virtual, menús de cuenta, herramientas de práctica y conexión con el backend mediante peticiones HTTP.", first_line=True)
    paragraph(doc, "El backend se implementó con FastAPI y expone rutas para salud del sistema, autenticación, chat, conversaciones, perfil de aprendizaje, biblioteca, solicitudes de conocimiento, herramientas educativas, retroalimentación y casos clínicos. La configuración del proyecto permite ejecutar el backend localmente con Uvicorn y servir el frontend desde la misma aplicación o como sitio estático conectado mediante una URL pública.", first_line=True)
    paragraph(doc, "La arquitectura de IA centraliza la selección de modelos. Groq se utiliza como proveedor principal para texto, razonamiento y visión; Cloudflare Workers AI funciona como respaldo cuando el proveedor principal falla por tiempo de espera, límite de uso o indisponibilidad; Gemini se reserva para la generación de imágenes educativas. Las respuestas médicas se clasifican por estado de verificación para que el sistema no presente como comprobado lo que no tiene evidencia documental.", first_line=True)

    paragraph(doc, "Etapas de desarrollo", bold=True, after=4)
    add_table(
        doc,
        ["Etapa", "Actividad", "Resultado obtenido", "Responsable"],
        [
            ("1", "Análisis de necesidad educativa y definición del alcance médico.", "Se delimitó MEDIA como tutor educativo, no diagnóstico.", "Equipo completo"),
            ("2", "Diseño del frontend web con chat, biblioteca, historial y paneles de herramientas.", "Interfaz funcional en HTML, CSS y JavaScript.", "Equipo frontend"),
            ("3", "Implementación del backend con FastAPI y modelos Pydantic.", "API organizada por módulos: chat, auth, tools, library, learning y feedback.", "Equipo backend"),
            ("4", "Integración de proveedores de IA y fallback.", "Groq, Cloudflare y Gemini separados según uso técnico.", "Equipo IA"),
            ("5", "Preparación de referencias médicas y RAG.", "Carpeta de PDFs, Markdown para RAG e ingesta a Supabase.", "Equipo de datos"),
            ("6", "Pruebas y ajustes.", "Pruebas automatizadas para contratos, rutas, recuperación y orquestación.", "Equipo completo"),
        ],
    )

    paragraph(doc, "Tecnologías y costos asociados", bold=True, after=4)
    add_table(
        doc,
        ["Recurso", "Uso dentro del proyecto", "Tipo de costo", "Costo estimado"],
        [
            ("Computadora de desarrollo", "Programación, pruebas locales y preparación del prototipo.", "Equipo disponible", "C$ 0"),
            ("Python 3.12 y FastAPI", "Backend, rutas API, servicios y validación.", "Software libre", "C$ 0"),
            ("HTML, CSS y JavaScript", "Interfaz web, chat, biblioteca y herramientas visuales.", "Software libre", "C$ 0"),
            ("Supabase/Postgres/pgvector", "Autenticación, datos, conversaciones, referencias y embeddings.", "Plan gratuito o escalable", "C$ 0 en prototipo"),
            ("Groq API", "Modelo principal para chat y razonamiento.", "Servicio API", "Variable según uso"),
            ("Cloudflare Workers AI", "Respaldo de texto y embeddings.", "Servicio API", "Variable según uso"),
            ("Gemini API", "Generación de imágenes educativas.", "Servicio API", "Variable según uso"),
            ("Vercel o servidor web", "Despliegue del frontend o aplicación.", "Plan gratuito o escalable", "C$ 0 en prototipo"),
        ],
    )

    paragraph(doc, "Análisis de desarrollo a pequeña y gran escala", bold=True, after=4)
    add_table(
        doc,
        ["Escala", "Aplicación", "Consideraciones técnicas"],
        [
            ("Pequeña escala", "Uso local o demostración para un grupo reducido de estudiantes.", "Puede ejecutarse con backend local, frontend estático, referencias locales y plan gratuito de servicios externos."),
            ("Mediana escala", "Uso institucional en una clase, laboratorio o facultad.", "Requiere backend público, control de usuarios, CORS configurado, base de datos estable y políticas de acceso a documentos."),
            ("Gran escala", "Plataforma educativa para múltiples cursos o instituciones.", "Necesita monitoreo, límites por usuario, seguridad reforzada, costos de IA controlados, alta disponibilidad e ingesta documental supervisada."),
        ],
    )

    heading(doc, "VI", "Impacto del proyecto")
    paragraph(doc, "MEDIA aporta valor porque convierte una necesidad académica en una solución tecnológica integrada. En lugar de depender de búsquedas aisladas o respuestas sin trazabilidad, el estudiante puede conversar con un tutor educativo, consultar una biblioteca de documentos, generar materiales de práctica y recibir respuestas clasificadas según el respaldo disponible.", first_line=True)
    paragraph(doc, "El proyecto ayuda a optimizar el tiempo de estudio, organizar temas médicos complejos y reforzar el aprendizaje mediante recursos como quizzes, flashcards, mapas mentales, presentaciones, informes e imágenes educativas. Además, promueve una cultura de uso responsable de la inteligencia artificial al separar conocimiento general no verificado de evidencia médica recuperada desde documentos.", first_line=True)
    paragraph(doc, "En el entorno real, MEDIA puede servir como prototipo para laboratorios de innovación educativa, ferias STEM, proyectos de software académico o iniciativas de apoyo al aprendizaje clínico. Su mayor impacto está en demostrar que la IA puede diseñarse con límites, transparencia y enfoque pedagógico.", first_line=True)

    heading(doc, "VII", "Conclusiones")
    paragraph(doc, "El proyecto MEDIA cumple con el objetivo de proponer una plataforma educativa capaz de apoyar el estudio médico mediante inteligencia artificial, biblioteca digital y herramientas de práctica. La solución integra frontend, backend, base de datos, proveedores de IA y recuperación documental en una arquitectura coherente.", first_line=True)
    paragraph(doc, "Los objetivos específicos se cumplen al contar con una interfaz de chat, rutas de backend organizadas, integración de modelos, manejo de usuarios, biblioteca de referencias y controles de verificación. La estructura del proyecto evidencia que no se trata solo de un chat, sino de un sistema diseñado para aprendizaje y consulta académica responsable.", first_line=True)
    paragraph(doc, "Como resultado, MEDIA representa una solución viable para el contexto educativo. Su versión prototipo puede ejecutarse localmente y su arquitectura permite crecer hacia una implementación institucional con mayor cantidad de usuarios, documentos, métricas y supervisión docente.", first_line=True)

    heading(doc, "VIII", "Recomendaciones")
    bullet(doc, "Completar la asignación formal del tutor y la carrera de cada participante antes de la entrega final impresa.")
    bullet(doc, "Realizar una demostración en vivo con una pregunta médica educativa, una consulta a la biblioteca y una herramienta de práctica.")
    bullet(doc, "Documentar capturas del frontend, el flujo de autenticación, el chat, la biblioteca y la generación de recursos educativos.")
    bullet(doc, "Configurar límites de uso, protección de claves API y monitoreo antes de un despliegue público.")
    bullet(doc, "Ampliar la base de referencias médicas con documentos revisados y clasificados por tema para mejorar la recuperación RAG.")
    bullet(doc, "Mantener visible el aviso de que MEDIA es una herramienta educativa y no sustituye atención médica profesional.")

    heading(doc, "IX", "Referencias bibliográficas")
    refs = [
        "FastAPI. (s. f.). FastAPI documentation. https://fastapi.tiangolo.com/",
        "Pydantic. (s. f.). Pydantic documentation. https://docs.pydantic.dev/",
        "Supabase. (s. f.). Supabase documentation. https://supabase.com/docs",
        "Groq. (s. f.). Groq API documentation. https://console.groq.com/docs",
        "Cloudflare. (s. f.). Workers AI documentation. https://developers.cloudflare.com/workers-ai/",
        "Google AI for Developers. (s. f.). Gemini API documentation. https://ai.google.dev/",
        "Lewis, P., Perez, E., Piktus, A., Petroni, F., Karpukhin, V., Goyal, N., et al. (2020). Retrieval-augmented generation for knowledge-intensive NLP tasks. Advances in Neural Information Processing Systems.",
        "OpenAI. (2024). Prompt engineering and model safety concepts. https://platform.openai.com/docs",
    ]
    for ref in refs:
        p = paragraph(doc, ref, after=4)
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.first_line_indent = Inches(-0.25)

    heading(doc, "X", "Anexos")
    paragraph(doc, "Anexo 1. Estructura principal del proyecto", bold=True, after=2)
    for line in [
        "index.html: estructura de la interfaz web.",
        "styles.css: diseño visual de la aplicación.",
        "app.js: lógica del chat, biblioteca, historial, autenticación y herramientas.",
        "backend/app/main.py: aplicación FastAPI y registro de rutas.",
        "backend/app/ai/: orquestación, proveedores, recuperación y verificación.",
        "backend/app/api/: endpoints de chat, auth, learning, library, tools y feedback.",
        "Referencias/: PDFs médicos y documentos Markdown preparados para RAG.",
    ]:
        bullet(doc, line)

    paragraph(doc, "Anexo 2. Rutas principales del backend", bold=True, after=2)
    add_table(
        doc,
        ["Ruta", "Función"],
        [
            ("GET /api/health", "Verifica el estado del backend."),
            ("POST /api/auth/register", "Registra usuarios."),
            ("POST /api/auth/login", "Inicia sesión y permite acceso al chat."),
            ("POST /api/chat", "Envía preguntas al tutor médico IA."),
            ("GET /api/library", "Consulta el catálogo de referencias."),
            ("POST /api/tools/quiz", "Genera quizzes educativos."),
            ("POST /api/tools/flashcards", "Genera tarjetas de estudio."),
            ("POST /api/tools/image", "Genera imágenes educativas con Gemini."),
        ],
    )

    paragraph(doc, "Anexo 3. Comandos de ejecución local", bold=True, after=2)
    for line in [
        "Backend: cd backend && uvicorn app.main:app --reload --host 127.0.0.1 --port 8000",
        "Frontend: python -m http.server 5500",
        "Pruebas: cd backend && pytest",
        "Ingesta de referencias: python backend/scripts/ingest_references.py --dry-run",
    ]:
        p = paragraph(doc, line, after=1, size=9.5)
        for run in p.runs:
            format_run(run, size=9.5, font="Courier New")

    doc.save(OUTPUT)


if __name__ == "__main__":
    build()
