from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = r"C:\Users\Maykel\Documents\MEDIA\Informe_Feria_STEM_2026_Riego_Automatizado.docx"
LOGO = r"C:\Users\Maykel\Documents\MEDIA\docx_work\template_media\image1.png"


participants = [
    "Maikel Leandro Martinez Flores",
    "Michael Yurem Rugama Montenegro",
    "Osmar Alejandro Castro Mercado",
    "Lleyton Antonio Mendez Montiel",
    "Helmut",
]


def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_borders = tc_pr.first_child_found_in("w:tcBorders")
    if tc_borders is None:
        tc_borders = OxmlElement("w:tcBorders")
        tc_pr.append(tc_borders)

    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        if edge in kwargs:
            edge_data = kwargs.get(edge)
            tag = "w:{}".format(edge)
            element = tc_borders.find(qn(tag))
            if element is None:
                element = OxmlElement(tag)
                tc_borders.append(element)
            for key in ["sz", "val", "color", "space"]:
                if key in edge_data:
                    element.set(qn("w:{}".format(key)), str(edge_data[key]))


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=120, start=120, bottom=120, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def format_run(run, bold=False, italic=False, size=12, font="Arial"):
    run.bold = bold
    run.italic = italic
    run.font.name = font
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0, 0, 0)
    run._element.rPr.rFonts.set(qn("w:ascii"), font)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), font)


def add_paragraph(doc, text="", bold=False, italic=False, align=None, space_after=6, first_line=False):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(space_after)
    if first_line:
        p.paragraph_format.first_line_indent = Inches(0.35)
    run = p.add_run(text)
    format_run(run, bold=bold, italic=italic)
    return p


def add_heading(doc, number, title):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.5
    r1 = p.add_run(f"{number}.\t")
    format_run(r1, bold=True)
    r2 = p.add_run(title)
    format_run(r2, bold=True)
    return p


def style_table(table, widths=None):
    border = {"val": "single", "sz": "6", "color": "D9D9D9", "space": "0"}
    tr_pr = table.rows[0]._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)
    for row_index, row in enumerate(table.rows):
        for col_index, cell in enumerate(row.cells):
            set_cell_border(cell, top=border, left=border, bottom=border, right=border)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            if widths:
                cell.width = widths[col_index]
            for p in cell.paragraphs:
                p.paragraph_format.line_spacing = 1.15
                p.paragraph_format.space_after = Pt(0)
                for run in p.runs:
                    format_run(run, size=10.5)
            if row_index == 0:
                shade_cell(cell, "1F4E79")
                for p in cell.paragraphs:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for run in p.runs:
                        run.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)
            elif row_index % 2 == 0:
                shade_cell(cell, "F3F6FA")


def add_bullet(doc, text):
    p = doc.add_paragraph(style=None)
    p.paragraph_format.left_indent = Inches(0.25)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run("- " + text)
    format_run(run)


def build():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(12)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")

    # Cover page.
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(8)
    p.add_run().add_picture(LOGO, width=Inches(2.25))

    p = add_paragraph(doc, "TU UNIVERSIDAD POR EXCELENCIA", bold=True, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=8)
    for run in p.runs:
        run.font.size = Pt(12)

    cover_box = doc.add_table(rows=1, cols=1)
    cover_box.autofit = False
    cover_box.columns[0].width = Inches(6.2)
    cell = cover_box.cell(0, 0)
    set_cell_border(cell, top={"val": "single", "sz": "18", "color": "000000", "space": "0"},
                    left={"val": "single", "sz": "18", "color": "000000", "space": "0"},
                    bottom={"val": "single", "sz": "18", "color": "000000", "space": "0"},
                    right={"val": "single", "sz": "18", "color": "000000", "space": "0"})
    set_cell_margins(cell, top=120, start=140, bottom=120, end=140)
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = cell.paragraphs[0].add_run("FACULTAD DE INGENIERÍAS Y ARQUITECTURA")
    format_run(r, bold=True, size=14, font="Arial")
    p2 = cell.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p2.add_run("Feria STEM 2026 - 4ta Edición")
    format_run(r, bold=True, size=13, font="Arial")

    add_paragraph(doc, "", space_after=8)
    add_paragraph(doc, "Nombre del proyecto:", bold=True, space_after=2)
    add_paragraph(doc, "Sistema de riego automatizado con Arduino y sensor de humedad", space_after=8)
    add_paragraph(doc, "Participantes:", bold=True, space_after=2)
    for name in participants:
        p = add_paragraph(doc, name, space_after=0)
        p.paragraph_format.line_spacing = 1.0
        for run in p.runs:
            run.font.size = Pt(10.5)
    add_paragraph(doc, "", space_after=6)
    add_paragraph(doc, "Tutor:", bold=True, space_after=2)
    add_paragraph(doc, "Nombre completo", space_after=0)
    p = add_paragraph(doc, "Fecha de presentación: Octubre, 2026", align=WD_ALIGN_PARAGRAPH.RIGHT, space_after=0)
    p.paragraph_format.space_before = Pt(4)

    doc.add_page_break()

    p = add_paragraph(doc, "Índice", bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=12)
    for item in [
        "I. Introducción",
        "II. Planteamiento de problema",
        "III. Justificación",
        "IV. Objetivos",
        "V. Desarrollo técnico y metodológico",
        "VI. Impacto del proyecto",
        "VII. Conclusiones",
        "VIII. Recomendaciones",
        "IX. Referencias bibliográficas",
        "X. Anexos",
    ]:
        add_paragraph(doc, item, space_after=2)

    doc.add_page_break()

    add_heading(doc, "I", "Introducción")
    add_paragraph(doc, "El proyecto consiste en el diseño y construcción de un sistema de riego automatizado para huertos escolares o áreas verdes pequeñas. La propuesta utiliza una placa Arduino, un sensor de humedad del suelo, una bomba de agua y un módulo de control que permite activar el riego únicamente cuando la tierra se encuentra por debajo de un nivel de humedad establecido.", first_line=True)
    add_paragraph(doc, "Su importancia se relaciona con el uso responsable del agua y con la necesidad de incorporar soluciones tecnológicas sencillas a problemas del entorno cotidiano. En muchos espacios educativos el riego se realiza de forma manual, sin mediciones y en horarios irregulares, lo que puede provocar desperdicio de agua o descuido de las plantas.", first_line=True)
    add_paragraph(doc, "La solución busca demostrar cómo la tecnología y la ingeniería pueden apoyar la toma de decisiones mediante sensores, programación y control automático. Además de resolver una necesidad práctica, el prototipo permite a los participantes aplicar conocimientos de electrónica, lógica de programación, diseño de circuitos, medición de variables y evaluación de resultados.", first_line=True)

    add_heading(doc, "II", "Planteamiento de problema")
    add_paragraph(doc, "En los huertos escolares y jardines pequeños, el riego suele depender de la disponibilidad de una persona y de una estimación visual del estado del suelo. Esta práctica puede generar dos problemas: exceso de riego, que desperdicia agua y puede afectar las raíces, o falta de riego, que disminuye el crecimiento de las plantas.", first_line=True)
    add_paragraph(doc, "La necesidad principal es contar con un mecanismo económico y fácil de operar que mida la humedad del suelo y active el suministro de agua solo cuando sea necesario. El proyecto atiende esta situación mediante un prototipo que automatiza el proceso y reduce la dependencia del riego manual.", first_line=True)

    add_heading(doc, "III", "Justificación")
    add_paragraph(doc, "El equipo decidió realizar este proyecto porque combina una problemática real con una solución tecnológica alcanzable para una feria STEM. El riego automatizado permite observar de manera directa cómo un sensor transforma una condición del ambiente en una acción concreta.", first_line=True)
    add_paragraph(doc, "El prototipo sirve para optimizar el uso del agua, mejorar el cuidado de plantas y promover hábitos de sostenibilidad en espacios educativos. Beneficia directamente a estudiantes, docentes y responsables de áreas verdes, porque facilita el mantenimiento de cultivos pequeños y permite registrar mejores prácticas de riego.", first_line=True)
    add_paragraph(doc, "También tiene valor formativo, ya que integra conocimientos de ingeniería, electrónica básica, programación y análisis de datos. La experiencia de construirlo permite comprender el ciclo completo de un proyecto tecnológico: identificación del problema, diseño, construcción, prueba, ajuste y presentación de resultados.", first_line=True)

    add_heading(doc, "IV", "Objetivos")
    add_paragraph(doc, "Objetivo general:", bold=True, space_after=2)
    add_paragraph(doc, "Diseñar y construir un prototipo de riego automatizado que utilice sensores de humedad y control electrónico para optimizar el suministro de agua en un huerto escolar o área verde pequeña.", first_line=True)
    add_paragraph(doc, "Objetivos específicos:", bold=True, space_after=2)
    add_bullet(doc, "Identificar los componentes electrónicos y mecánicos necesarios para medir la humedad del suelo y activar una bomba de agua.")
    add_bullet(doc, "Programar una placa Arduino para comparar la lectura del sensor con un valor de referencia y ejecutar el riego de forma automática.")
    add_bullet(doc, "Construir y probar el circuito del prototipo verificando su funcionamiento en condiciones de suelo seco y húmedo.")
    add_bullet(doc, "Analizar la viabilidad del sistema a pequeña escala y proponer mejoras para su aplicación en espacios más amplios.")

    add_heading(doc, "V", "Desarrollo técnico y metodológico")
    add_paragraph(doc, "El desarrollo inició con la observación del problema y la definición de los requerimientos principales: medir humedad, controlar una bomba, evitar activaciones innecesarias y mantener el sistema simple para facilitar su explicación durante la feria. Después se diseñó un circuito básico con Arduino como unidad de control, sensor de humedad como entrada y bomba de agua como salida.", first_line=True)
    add_paragraph(doc, "El procedimiento se organizó en etapas para asegurar que cada parte del prototipo funcionara antes de integrar el sistema completo.", first_line=True)

    table = doc.add_table(rows=1, cols=4)
    table.rows[0].cells[0].text = "Etapa"
    table.rows[0].cells[1].text = "Actividad"
    table.rows[0].cells[2].text = "Resultado esperado"
    table.rows[0].cells[3].text = "Responsable"
    rows = [
        ("1", "Diagnóstico del problema y definición del alcance.", "Necesidad de automatizar el riego en un espacio pequeño.", "Equipo completo"),
        ("2", "Selección de componentes y diseño del circuito.", "Lista de materiales y diagrama de conexión.", "Área técnica"),
        ("3", "Programación del Arduino.", "Código capaz de leer humedad y activar la bomba.", "Área de programación"),
        ("4", "Construcción y pruebas del prototipo.", "Sistema funcional con riego automático.", "Equipo completo"),
        ("5", "Ajustes y preparación de la exposición.", "Prototipo estable y explicación clara para la feria.", "Equipo completo"),
    ]
    for row in rows:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = value
    style_table(table, widths=[Inches(0.6), Inches(2.0), Inches(2.25), Inches(1.3)])

    add_paragraph(doc, "Presupuesto invertido", bold=True, space_after=4)
    table = doc.add_table(rows=1, cols=5)
    headers = ["Cantidad", "Material", "Descripción de uso", "Costo unitario", "Subtotal"]
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = h
    materials = [
        ("1", "Arduino Uno compatible", "Control principal del sistema", "C$ 450", "C$ 450"),
        ("1", "Sensor de humedad capacitivo", "Medición de humedad del suelo", "C$ 180", "C$ 180"),
        ("1", "Módulo relé 5V", "Activación segura de la bomba", "C$ 90", "C$ 90"),
        ("1", "Bomba de agua 5V", "Suministro de agua al cultivo", "C$ 220", "C$ 220"),
        ("1", "Manguera y recipiente", "Conducción y reserva de agua", "C$ 160", "C$ 160"),
        ("1", "Protoboard y cables", "Conexiones eléctricas de prueba", "C$ 180", "C$ 180"),
        ("1", "Fuente de alimentación", "Energía para el circuito", "C$ 250", "C$ 250"),
        ("1", "Base o estructura", "Soporte del prototipo", "C$ 150", "C$ 150"),
        ("", "Total estimado", "", "", "C$ 1,680"),
    ]
    for row in materials:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = value
    style_table(table, widths=[Inches(0.65), Inches(1.45), Inches(2.0), Inches(1.0), Inches(1.0)])

    add_paragraph(doc, "Análisis de escalabilidad", bold=True, space_after=4)
    table = doc.add_table(rows=1, cols=3)
    for i, h in enumerate(["Escala", "Aplicación", "Consideraciones técnicas"]):
        table.rows[0].cells[i].text = h
    for row in [
        ("Pequeña escala", "Maceteras, huertos escolares o jardines de demostración.", "Un sensor y una bomba pequeña son suficientes; el mantenimiento es sencillo y el costo se mantiene bajo."),
        ("Mediana escala", "Varias camas de cultivo o áreas verdes de una institución.", "Se requieren más sensores, válvulas y una fuente de energía con mayor capacidad."),
        ("Gran escala", "Sistemas agrícolas o viveros con varias zonas de riego.", "Conviene integrar panel solar, almacenamiento de datos y control por sectores para mejorar eficiencia."),
    ]:
        cells = table.add_row().cells
        for index, value in enumerate(row):
            cells[index].text = value
    style_table(table, widths=[Inches(1.25), Inches(2.15), Inches(3.1)])

    add_heading(doc, "VI", "Impacto del proyecto")
    add_paragraph(doc, "El proyecto aporta una solución práctica porque permite regar las plantas de acuerdo con una medición y no únicamente por rutina. Esto ayuda a ahorrar agua, reduce el tiempo dedicado al riego manual y disminuye el riesgo de descuido en días de alta temperatura o poca supervisión.", first_line=True)
    add_paragraph(doc, "En el entorno educativo, el prototipo funciona como herramienta de aprendizaje. Los estudiantes pueden observar la relación entre una variable ambiental, una lectura electrónica y una respuesta automatizada. También promueve el pensamiento crítico, ya que obliga a ajustar valores, comparar resultados y justificar decisiones técnicas.", first_line=True)

    add_heading(doc, "VII", "Conclusiones")
    add_paragraph(doc, "El prototipo demuestra que es posible construir un sistema de riego automatizado con componentes accesibles y conocimientos básicos de programación y electrónica. El objetivo general se cumple al integrar sensor, Arduino, relé y bomba en un circuito funcional orientado al ahorro de agua.", first_line=True)
    add_paragraph(doc, "Los objetivos específicos se alcanzan porque el equipo identifica los materiales necesarios, programa la lógica de activación, realiza pruebas con diferentes niveles de humedad y analiza posibilidades de mejora. La propuesta confirma que la automatización puede resolver problemas cotidianos cuando se diseña a partir de una necesidad concreta.", first_line=True)
    add_paragraph(doc, "Como resultado, el proyecto tiene potencial para aplicarse en huertos escolares y demostraciones educativas. Su valor principal está en mostrar una solución sencilla, replicable y relacionada con el uso responsable de recursos naturales.", first_line=True)

    add_heading(doc, "VIII", "Recomendaciones")
    add_bullet(doc, "Calibrar el sensor antes de cada demostración para definir correctamente los valores de suelo seco y húmedo.")
    add_bullet(doc, "Proteger las conexiones eléctricas del contacto directo con agua para evitar fallas o cortocircuitos.")
    add_bullet(doc, "Incorporar una pantalla LCD o indicadores LED que muestren el estado del sistema durante la exposición.")
    add_bullet(doc, "Evaluar el uso de energía solar si el prototipo se implementa en exteriores por periodos prolongados.")
    add_bullet(doc, "Registrar mediciones en diferentes momentos del día para mejorar la precisión del umbral de riego.")

    add_heading(doc, "IX", "Referencias bibliográficas")
    refs = [
        "Arduino. (s. f.). Arduino documentation. https://docs.arduino.cc/",
        "Monk, S. (2016). Programming Arduino: Getting Started with Sketches (2nd ed.). McGraw-Hill Education.",
        "Organización de las Naciones Unidas para la Alimentación y la Agricultura. (s. f.). Water management. https://www.fao.org/water/",
    ]
    for ref in refs:
        p = add_paragraph(doc, ref, space_after=4)
        p.paragraph_format.first_line_indent = Inches(-0.25)
        p.paragraph_format.left_indent = Inches(0.25)

    add_heading(doc, "X", "Anexos")
    add_paragraph(doc, "Anexo 1. Código fuente básico", bold=True, space_after=3)
    code_lines = [
        "const int sensor = A0;",
        "const int rele = 8;",
        "const int limite = 500;",
        "",
        "void setup() {",
        "  pinMode(rele, OUTPUT);",
        "  Serial.begin(9600);",
        "}",
        "",
        "void loop() {",
        "  int humedad = analogRead(sensor);",
        "  if (humedad > limite) {",
        "    digitalWrite(rele, HIGH);",
        "  } else {",
        "    digitalWrite(rele, LOW);",
        "  }",
        "  delay(1000);",
        "}",
    ]
    for line in code_lines:
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(line)
        format_run(r, size=9.5, font="Courier New")

    add_paragraph(doc, "Anexo 2. Manual de usuario", bold=True, space_after=3)
    add_bullet(doc, "Conectar la fuente de alimentación del prototipo.")
    add_bullet(doc, "Colocar el sensor en la tierra sin forzar los cables.")
    add_bullet(doc, "Llenar el recipiente de agua y verificar que la manguera esté conectada.")
    add_bullet(doc, "Observar si la bomba se activa cuando el suelo está seco y se detiene cuando alcanza humedad suficiente.")
    add_paragraph(doc, "Anexo 3. Diagramas y evidencias", bold=True, space_after=3)
    add_paragraph(doc, "Se recomienda adjuntar fotografías del prototipo, diagrama del circuito, capturas del código y enlace a un video de demostración.", first_line=True)

    doc.save(OUTPUT)


if __name__ == "__main__":
    build()
